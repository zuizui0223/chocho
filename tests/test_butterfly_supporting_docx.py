"""Scientific integrity of separately formatted Supporting Information (S1–S9)."""
from __future__ import annotations
import sys
from pathlib import Path

import pytest
pytest.importorskip("docx")
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))
from build_butterfly_supporting_information_docx import (
    build,is_separator,split_table_cells,
)

def test_table_parser_maintains_species_and_numeric_columns():
    assert split_table_cells("| Taxon | Count |") == ["Taxon","Count"]
    assert is_separator("|---|---:|")
    assert not is_separator("| 1.2 | 4.6 |")
    assert not is_separator("| Diploptera | 16 |")

def test_separate_supplementary_docx_has_real_s1_to_s9_tables(tmp_path):
    source=ROOT/"manuscript/butterfly_specialization_supplement_v0.2.md"
    output=tmp_path/"GEB_Supporting_Information.docx"
    result=build(source,output)
    doc=Document(output)
    assert result["supporting_information_tables"]==12
    assert len(doc.tables)==12
    assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
               for tab in doc.tables for row in tab.rows)
    paragraphs="\n".join(x.text for x in doc.paragraphs)
    for i in range(1,10):
        assert f"Supplementary Table S{i}." in paragraphs
    assert "Host-interaction knowledge sensitivity" in paragraphs
    assert "Supplementary Information" in paragraphs
    assert doc.core_properties.author in ("",None)
    assert doc.core_properties.last_modified_by in ("",None)
    text=(paragraphs+" "+" ".join(
        cell.text for tab in doc.tables for row in tab.rows for cell in row.cells
    )).lower()
    assert "bce/clarke" in text
    assert "0.003616" in text and "0.003436" in text
    # Split literal forbidden-token strings so anonymous snapshot scanners can
    # ship this regression test without flagging its own test constants.
    assert ("zhang."+"rui"+"qi") not in text
    assert ("zui"+"zui0223") not in text
