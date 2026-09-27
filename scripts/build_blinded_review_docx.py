#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import etree


FIGURES = {
    1: "Figure1_taxonomic_vs_geographic_specialization.png",
    2: "Figure2_anthropogenic_resource_expansion.png",
    3: "Figure3_host_contribution_architecture.png",
    4: "Figure4_within_family_specialization_hierarchy.png",
    5: "Figure5_independent_climate_filtering.png",
}


def set_run_font(run, name: str, size: float | None = None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline_markdown(paragraph, text: str, *, size: float = 11.0):
    token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\x60[^\x60]+\x60)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            r = paragraph.add_run(text[pos:match.start()])
            set_run_font(r, "Times New Roman", size)
        token = match.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            set_run_font(r, "Times New Roman", size, bold=True)
        elif token.startswith("*"):
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, "Times New Roman", size, italic=True)
        else:
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, "Courier New", size - 1)
        pos = match.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, "Times New Roman", size)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, "Times New Roman", 9)


def add_continuous_line_numbering(section):
    sectpr = section._sectPr
    for old in sectpr.findall(qn("w:lnNumType")):
        sectpr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "continuous")
    sectpr.append(ln)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(28)
    section.right_margin = Mm(25)
    add_continuous_line_numbering(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for name, size in [("Heading 1", 13), ("Heading 2", 12)]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "ManuscriptTitle" not in styles:
        title = styles.add_style("ManuscriptTitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title = styles["ManuscriptTitle"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(14)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True

    if "BlockQuote" not in styles:
        quote = styles.add_style("BlockQuote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["BlockQuote"]
    quote.font.name = "Times New Roman"
    quote.font.size = Pt(11)
    quote.font.italic = True
    quote.font.color.rgb = RGBColor(0, 0, 0)
    quote.paragraph_format.left_indent = Mm(8)
    quote.paragraph_format.right_indent = Mm(4)
    quote.paragraph_format.line_spacing = 1.5
    quote.paragraph_format.space_after = Pt(6)

    if "CodeBlock" not in styles:
        code = styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["CodeBlock"]
    code.font.name = "Courier New"
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Mm(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(4)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].clear()
    add_page_number(footer.paragraphs[0])

    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = "Double-anonymous review manuscript"
    props.keywords = ""
    props.category = ""
    props.title = "Blinded review manuscript"


def add_body_paragraph(doc: Document, text: str, *, references: bool = False):
    p = doc.add_paragraph()
    if references:
        p.paragraph_format.left_indent = Mm(6)
        p.paragraph_format.first_line_indent = Mm(-6)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        add_inline_markdown(p, text, size=10.5)
    else:
        add_inline_markdown(p, text)
    return p


def build_docx(markdown_path: Path, figures_dir: Path, output_path: Path, review_link: str | None = None):
    text = markdown_path.read_text(encoding="utf-8")
    if review_link:
        text = text.replace("[ANONYMIZED REVIEW LINK]", review_link)
    doc = Document()
    configure_document(doc)

    in_code = False
    in_references = False
    in_figure_legends = False
    code_lines: list[str] = []

    lines = text.splitlines()
    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph("\n".join(code_lines), style="CodeBlock")
                for run in p.runs:
                    set_run_font(run, "Courier New", 9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue
        if line.strip() == "---":
            continue

        if line.startswith("# "):
            p = doc.paragraphs[0]
            p.style = "ManuscriptTitle"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p, line[2:].strip(), size=14)
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            in_references = heading.startswith("References")
            in_figure_legends = heading == "Figure legends"
            display_heading = "References" if in_references else heading
            p = doc.add_paragraph(display_heading, style="Heading 1")
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip(), style="Heading 2")
            continue

        if line.startswith("**Running title:**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p, line)
            continue

        if line.startswith("> "):
            p = doc.add_paragraph(style="BlockQuote")
            add_inline_markdown(p, line[2:].strip())
            continue

        figure_match = re.match(r"^\*\*Figure\s+([1-5])\.", line)
        if in_figure_legends and figure_match:
            number = int(figure_match.group(1))
            if number > 1:
                doc.add_page_break()
            caption = add_body_paragraph(doc, line)
            caption.paragraph_format.keep_with_next = True
            image_path = figures_dir / FIGURES[number]
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_together = True
            run = p.add_run()
            run.add_picture(str(image_path), width=Inches(6.05))
            continue

        if line.startswith("- "):
            item = line[2:].strip()
            if in_references:
                add_body_paragraph(doc, item, references=True)
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_inline_markdown(p, item)
            continue

        add_body_paragraph(doc, line, references=in_references)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    scrub_docx(output_path)


def scrub_docx(path: Path):
    ns_core = {
        "dc": "http://purl.org/dc/elements/1.1/",
        "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
    }
    with zipfile.ZipFile(path, "r") as src:
        entries = [(info, src.read(info.filename)) for info in src.infolist()]

    tmp = path.with_suffix(".scrubbed.tmp.docx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in entries:
            name = info.filename
            if name == "docProps/custom.xml":
                continue
            if name == "docProps/core.xml":
                root = etree.fromstring(data)
                for xp in ("//dc:creator", "//cp:lastModifiedBy"):
                    for node in root.xpath(xp, namespaces=ns_core):
                        node.text = ""
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif name.startswith("word/") and name.endswith(".xml"):
                root = etree.fromstring(data)
                for elem in root.iter():
                    for attr in list(elem.attrib):
                        local = etree.QName(attr).localname
                        if local.startswith("rsid"):
                            del elem.attrib[attr]
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info, data)
    tmp.replace(path)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input-md", type=Path, required=True)
    ap.add_argument("--figures-dir", type=Path, required=True)
    ap.add_argument("--output", type=Path, required=True)
    ap.add_argument("--review-link", default=None)
    args = ap.parse_args()
    build_docx(args.input_md, args.figures_dir, args.output, args.review_link)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
