#!/usr/bin/env python3
"""Source availability audit for Clarke 2024 independently compiled foodplant records.

No ecological hypothesis fitted. Fail closed if neither the original Dryad CSV
nor the journal-supplied Appendix S1 can be retrieved with verifiable provenance.
"""
from __future__ import annotations
import argparse,csv,hashlib,io,json,time,urllib.request
from pathlib import Path
import zipfile

SOURCE_CANDIDATES = [
 ("original_dryad_csv_via_files", "https://datadryad.org/api/v2/files/2850981/download", "csv"),
 ("original_dryad_csv_stream", "https://datadryad.org/downloads/file_stream/2850981", "csv"),
 ("publisher_appendix_s1", "https://pmc.ncbi.nlm.nih.gov/articles/instance/10771928/bin/ECE3-14-e10834-s005.xlsx", "xlsx"),
]
SCHEMA_TITLE="A checklist of European butterfly larval foodplants"

def get(url):
    req=urllib.request.Request(url,headers={
        "User-Agent":"chocho-source-audit/1.0 (academic reproducibility, github.com/zuizui0223/chocho)",
        "Accept":"text/csv,application/zip,application/octet-stream,*/*"})
    with urllib.request.urlopen(req,timeout=35) as rsp:return rsp.read()

def examine(raw, expected):
    if expected=="csv":
        if not raw or b"\x00" in raw[:1000]:raise RuntimeError("not CSV")
        text=raw.decode("utf-8-sig")
        rd=csv.reader(io.StringIO(text))
        header=next(rd)
        if "butterfly_id" not in header or "status" not in header or "powo_id" not in header:
            raise RuntimeError("unexpected Clarke larval_host_plant.csv schema")
        count=sum(1 for _ in rd)
        if count<1000:raise RuntimeError("implausibly small compiled host table")
        return {"kind":"source_table_csv","columns":header,"rows":count}
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if not any(x.endswith("xl/workbook.xml") for x in z.namelist()):
            raise RuntimeError("XLSX archive workbook content missing")
    import openpyxl
    w=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    sheets=[]
    for sh in w:
        head=[]
        for row in sh.iter_rows(min_row=1,max_row=min(4,sh.max_row or 4),values_only=True):
            head.append([str(x)[:80] if x is not None else "" for x in row[:20]])
        sheets.append({"name":sh.title,"estimated_rows":sh.max_row,
                       "estimated_columns":sh.max_column,"head":head})
    return {"kind":"publisher_original_appendix","sheets":sheets}

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--protocol",required=True,type=Path)
    p.add_argument("--output",required=True,type=Path)
    args=p.parse_args()
    contract=json.loads(args.protocol.read_text(encoding="utf-8"))
    if contract["schema"]!="chocho_clarke2024_heldout_host_evidence_gate_v0.1":
        raise RuntimeError("unexpected source protocol")
    result={"schema":"chocho_clarke2024_source_feasibility_v0.1",
            "original_publication":SCHEMA_TITLE,"data_doi":contract["source"]["data"],
            "source_outcomes":[],"usable_original_source":None}
    for label,url,kind in SOURCE_CANDIDATES:
        attempts=[]
        source=None
        for rep in range(2):
            try:
                raw=get(url)
                content=examine(raw,kind)
                source={"label":label,"url":url,"sha256":hashlib.sha256(raw).hexdigest(),
                        "bytes":len(raw),"structure":content}
                break
            except Exception as e:
                attempts.append({"attempt":rep+1,"error":type(e).__name__+":"+str(e)[:200]})
                if rep==0:time.sleep(1)
        result["source_outcomes"].append(
            {"candidate":label,"verified":bool(source),"source":source,"failures":attempts})
        if source is not None and result["usable_original_source"] is None:
            result["usable_original_source"]=source
    result["decision"]="SOURCE_AVAILABLE_SCHEMA_ONLY" if result["usable_original_source"] else "HOLD_SOURCE_UNAVAILABLE"
    result["warning"]="A publisher supplement and raw Dryad database table may have DIFFERENT fields; do not treat the appendix as a byte-identical copy or infer independent biological observations from literature compilation."
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps({"decision":result["decision"],"candidates":[{
         "candidate":x["candidate"],"verified":x["verified"],
         "source":x["source"],"failures":x["failures"]} for x in result["source_outcomes"]]},indent=2),flush=True)
    # Source gate failure is a scientific HOLD, not silently a successful data analysis.
    if result["usable_original_source"] is None:
        raise SystemExit("Original Clarke source could not be acquired; independent biological analysis held")

if __name__=="__main__":
    main()
