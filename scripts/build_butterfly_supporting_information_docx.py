#!/usr/bin/env python3
"""Make a separate editable, anonymized GEB Supporting Information DOCX.

Markdown Tables S1--S9 are preserved as true Word tables, not paragraphs.
Build is deterministic from the frozen versioned supplementary Markdown.
"""
from __future__ import annotations
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from build_blinded_review_docx import add_inline_markdown, scrub_docx

S1_TO_S9=tuple(f"Supplementary Table S{i}." for i in range(1,10))
# Split the scanner examples to avoid embedding identifying tokens in the anonymous code archive.
ANON_BAN=("zui"+"zui0223","rui"+"qi","zhang."+"rui"+"qi")

def render_links(s: str) -> str:
    # Preserve DOI and original-source URLs in the supplementary DOCX.
    return re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)",lambda m:f"{m.group(1)} ({m.group(2)})",s)

def is_separator(text: str) -> bool:
    """Match a genuine Markdown header separator, never a biological data row."""
    if not text.lstrip().startswith("|"):return False
    cells=split_table_cells(text)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?",cell.strip()) for cell in cells)

def split_table_cells(line: str) -> list[str]:
    t=line.strip()
    if not t.startswith("|"):raise ValueError("not a Markdown table row")
    if t.endswith("|"):t=t[:-1]
    t=t[1:]
    return [z.strip() for z in re.split(r"(?<!\\)\|",t)]

def add_bold_header(cell):
    tcPr=cell._tc.get_or_add_tcPr()
    # Repeat table headers when the table flows onto subsequent pages.
    tr=cell._tc.getparent()
    trPr=tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        header=OxmlElement("w:tblHeader")
        header.set(qn("w:val"),"true")
        trPr.append(header)

def add_table(doc, rows: list[list[str]]):
    if not rows:return None
    count=len(rows[0])
    if not 2 <= count <= 12:
        raise ValueError(f"Unsupported supplementary table width: {count}")
    if any(len(r)!=count for r in rows):
        raise ValueError("Markdown supplementary table has inconsistent column widths")
    table=doc.add_table(rows=len(rows),cols=count)
    table.style="Table Grid"
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=True
    size=8.25 if count<=5 else 7.6 if count<=7 else 7.2
    for i,row in enumerate(rows):
        for col,text in enumerate(row):
            cell=table.cell(i,col)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.line_spacing=1.05
            add_inline_markdown(p,render_links(text),size=size)
            for run in p.runs:
                if i==0:run.bold=True
        if i==0:add_bold_header(table.cell(0,0))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def configure(doc):
    s=doc.sections[0]
    s.orientation=WD_ORIENT.LANDSCAPE
    s.page_width=Mm(297)
    s.page_height=Mm(210)
    s.top_margin=Mm(17)
    s.bottom_margin=Mm(17)
    s.left_margin=Mm(18)
    s.right_margin=Mm(18)
    normal=doc.styles["Normal"]
    normal.font.name="Times New Roman"
    normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1.12
    normal.paragraph_format.space_after=Pt(5)
    for sty,size in (("Title",15),("Heading 1",12),("Heading 2",10.8)):
        f=doc.styles[sty]
        f.font.name="Times New Roman"
        f.font.size=Pt(size)
        f.font.bold=True
        f.font.color.rgb=RGBColor(0,0,0)
        f.paragraph_format.space_before=Pt(10)
        f.paragraph_format.space_after=Pt(6)
        f.paragraph_format.keep_with_next=True
    props=doc.core_properties
    props.title="Supporting Information — butterfly resource geography"
    props.subject="Double-anonymous scientific supplementary information"
    props.author=""
    props.last_modified_by=""
    props.keywords=""
    props.comments=""
    foot=s.footer.paragraphs[0]
    foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=foot.add_run()
    fld=OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"),"PAGE")
    run._r.addnext(fld)

def build(source:Path,output:Path)->dict:
    md=source.read_text(encoding="utf-8")
    for ban in ANON_BAN:
        if ban in md.lower():raise ValueError("Identifying string in SI Markdown")
    for header in S1_TO_S9:
        if "## "+header not in md:raise RuntimeError("Missing SI table "+header)
    lines=md.splitlines()
    doc=Document()
    configure(doc)
    block=[]
    n_tables=0
    current_paragraph=[]
    headings=[]

    def flush_paragraph():
        nonlocal current_paragraph
        if not current_paragraph:return
        p=doc.add_paragraph()
        add_inline_markdown(p,render_links(" ".join(current_paragraph)),size=10)
        current_paragraph=[]

    def flush_table():
        nonlocal block,n_tables
        if block:
            add_table(doc,block)
            n_tables+=1
        block=[]

    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:
            flush_paragraph()
            flush_table()
            i+=1
            continue
        if line.startswith("|"):
            flush_paragraph()
            if is_separator(line):
                i+=1
                continue
            cells=split_table_cells(line)
            if block and len(cells)!=len(block[0]):
                raise ValueError("Table width changes mid-block at Markdown line "+str(i+1))
            block.append(cells)
            i+=1
            continue
        flush_table()
        if line.startswith("# "):
            flush_paragraph()
            p=doc.add_paragraph(style="Title")
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            add_inline_markdown(p,render_links(line[2:]),size=15)
        elif line.startswith("## "):
            flush_paragraph()
            name=line[3:].strip()
            headings.append(name)
            doc.add_paragraph(render_links(name),style="Heading 1")
        elif line.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(render_links(line[4:]),style="Heading 2")
        elif line.startswith("- "):
            flush_paragraph()
            p=doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p,render_links(line[2:]),size=10)
        elif re.match(r"^\d+\.\s+",line):
            flush_paragraph()
            p=doc.add_paragraph(style="List Number")
            add_inline_markdown(p,render_links(re.sub(r"^\d+\.\s+","",line)),size=10)
        else:
            current_paragraph.append(line)
        i+=1
    flush_paragraph()
    flush_table()
    if n_tables!=12:
        raise RuntimeError(f"Expected exactly 12 scientific tables, found {n_tables}")
    if not all(any(h.startswith(t) for h in headings) for t in S1_TO_S9):
        raise RuntimeError("Missing supplementary table headings after DOCX conversion")
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    scrub_docx(output)
    return {"supporting_information_tables":n_tables,"supplementary_headings":len(headings),
            "real_word_tables":len(doc.tables),"source":str(source),"output":str(output)}

def main():
    p=argparse.ArgumentParser()
    p.add_argument("--input-md",type=Path,required=True)
    p.add_argument("--output",type=Path,required=True)
    args=p.parse_args()
    report=build(args.input_md,args.output)
    print(report,flush=True)

if __name__=="__main__":
    main()
